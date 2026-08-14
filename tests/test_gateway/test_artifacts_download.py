from __future__ import annotations

import hashlib
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from openstarry_code.artifacts import ArtifactStore


class _FakeSessionManager:
    def __init__(self, session_id: str) -> None:
        self.session_id = session_id

    async def get_session(self, session_key: str) -> object | None:
        if session_key == "agent:main:webchat:ok":
            return SimpleNamespace(session_id=self.session_id)
        return None


def _app(tmp_path: Path, *, auth_mode: str = "token", host: str = "127.0.0.1"):
    pytest.importorskip("starlette.testclient")
    from starlette.applications import Starlette

    from openstarry_code.gateway.artifacts import register_artifact_routes
    from openstarry_code.gateway.config import AttachmentsConfig, AuthConfig, GatewayConfig
    from openstarry_code.gateway.middleware import AuthMiddleware

    config = GatewayConfig(
        host=host,
        auth=AuthConfig(mode=auth_mode, token="secret"),
        attachments=AttachmentsConfig(media_root=str(tmp_path)),
    )
    app = Starlette(debug=False)
    register_artifact_routes(
        app,
        config=config,
        session_manager=_FakeSessionManager("session-1"),
    )
    app.add_middleware(AuthMiddleware, config=config)
    return app


def _publish(
    tmp_path: Path,
    *,
    payload: bytes = b"hello artifact",
    name: str = "report final.txt",
    mime: str = "text/plain",
):
    return ArtifactStore(tmp_path).publish_bytes(
        payload,
        session_id="session-1",
        session_key="agent:main:webchat:ok",
        name=name,
        mime=mime,
        source="publish_artifact",
    )


def test_artifact_download_requires_auth_and_session_scope(tmp_path: Path) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    ref = _publish(tmp_path)

    with TestClient(_app(tmp_path)) as client:
        unauthenticated = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok"
        )
        wrong_session = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:other",
            headers={"Authorization": "Bearer secret"},
        )
        missing_session = client.get(
            f"/api/v1/artifacts/{ref.id}",
            headers={"Authorization": "Bearer secret"},
        )

    assert unauthenticated.status_code == 401
    assert wrong_session.status_code == 404
    assert missing_session.status_code == 404


def test_artifact_download_serves_file_response_headers_and_ranges(tmp_path: Path) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    ref = _publish(tmp_path)

    with TestClient(_app(tmp_path)) as client:
        response = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret"},
        )
        ranged = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret", "Range": "bytes=0-4"},
        )

    assert response.status_code == 200
    assert response.content == b"hello artifact"
    assert response.headers["content-type"].startswith("text/plain")
    assert "attachment" in response.headers["content-disposition"]
    assert "report%20final.txt" in response.headers["content-disposition"]
    assert ranged.status_code == 206
    assert ranged.content == b"hello"


def test_artifact_download_preserves_valid_pptx_bytes(tmp_path: Path) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "下载后仍可读取"
    output = io.BytesIO()
    presentation.save(output)
    payload = output.getvalue()
    ref = _publish(
        tmp_path,
        payload=payload,
        name="课程总结.pptx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
    )

    with TestClient(_app(tmp_path)) as client:
        response = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret"},
        )

    assert response.status_code == 200
    assert response.content == payload
    assert ref.sha256 == hashlib.sha256(response.content).hexdigest()
    assert response.headers["content-type"].startswith(ref.mime)
    assert "%E8%AF%BE%E7%A8%8B%E6%80%BB%E7%BB%93.pptx" in response.headers[
        "content-disposition"
    ]
    downloaded = Presentation(io.BytesIO(response.content))
    assert downloaded.slides[0].shapes.title is not None
    assert downloaded.slides[0].shapes.title.text == "下载后仍可读取"


def test_artifact_download_keeps_historical_invalid_pptx_available(tmp_path: Path) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    legacy_payload = b"historically stored bytes that are not OOXML"
    ref = _publish(
        tmp_path,
        payload=legacy_payload,
        name="legacy.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    with TestClient(_app(tmp_path)) as client:
        response = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret"},
        )

    assert response.status_code == 200
    assert response.content == legacy_payload


def test_artifact_download_reports_not_found_and_integrity_errors(tmp_path: Path) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    ref = _publish(tmp_path)
    ArtifactStore(tmp_path).path_for(ref).write_bytes(b"tampered")

    with TestClient(_app(tmp_path)) as client:
        missing = client.get(
            "/api/v1/artifacts/missing?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret"},
        )
        mismatch = client.get(
            f"/api/v1/artifacts/{ref.id}?sessionKey=agent:main:webchat:ok",
            headers={"Authorization": "Bearer secret"},
        )

    assert missing.status_code == 404
    assert mismatch.status_code == 409


def test_artifact_native_open_owner_opens_html_copy(tmp_path: Path, monkeypatch) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    from openstarry_code.gateway import artifacts as artifact_routes

    ref = _publish(
        tmp_path,
        payload=b"<!doctype html><title>ok</title>",
        name="report.html",
        mime="text/html",
    )
    opened: list[Path] = []
    monkeypatch.setattr(artifact_routes.tempfile, "gettempdir", lambda: str(tmp_path / "tmp"))
    monkeypatch.setattr(
        artifact_routes,
        "_open_path_with_default_app",
        lambda path: opened.append(Path(path)) or None,
    )

    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        response = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={
                "Authorization": "Bearer secret",
                "x-opensquilla-session-key": "agent:main:webchat:ok",
            },
        )

    assert response.status_code == 202
    assert response.json() == {"ok": True, "status": "accepted"}
    assert len(opened) == 1
    assert opened[0].name.endswith("-report.html")
    assert opened[0].read_bytes() == b"<!doctype html><title>ok</title>"


def test_artifact_open_cache_dir_skips_posix_mode_bits_on_windows(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from openstarry_code.gateway import artifacts as artifact_routes

    temp_root = tmp_path / "tmp"
    cache_root = temp_root / "opensquilla-artifacts"
    cache_root.mkdir(parents=True)
    cache_root.chmod(0o777)
    monkeypatch.setattr(artifact_routes.tempfile, "gettempdir", lambda: str(temp_root))
    monkeypatch.setattr(artifact_routes.sys, "platform", "win32")

    assert artifact_routes._artifact_open_cache_dir() == cache_root


def test_artifact_native_open_requires_auth_and_owner(tmp_path: Path, monkeypatch) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    from openstarry_code.gateway import artifacts as artifact_routes

    ref = _publish(tmp_path, payload=b"<html></html>", name="page.html", mime="text/html")
    opened: list[Path] = []
    monkeypatch.setattr(
        artifact_routes,
        "_open_path_with_default_app",
        lambda path: opened.append(Path(path)) or None,
    )

    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        unauthenticated = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={"x-opensquilla-session-key": "agent:main:webchat:ok"},
        )
    with TestClient(
        _app(tmp_path, auth_mode="none", host="0.0.0.0"),
        client=("127.0.0.1", 50000),
    ) as client:
        non_owner = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={"x-opensquilla-session-key": "agent:main:webchat:ok"},
        )

    assert unauthenticated.status_code == 401
    assert non_owner.status_code == 403
    assert non_owner.json()["code"] == "OWNER_REQUIRED"
    assert opened == []


def test_artifact_native_open_requires_session_scope_and_integrity(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    from openstarry_code.gateway import artifacts as artifact_routes

    ref = _publish(tmp_path, payload=b"<html></html>", name="page.html", mime="text/html")
    monkeypatch.setattr(artifact_routes.tempfile, "gettempdir", lambda: str(tmp_path / "tmp"))
    monkeypatch.setattr(artifact_routes, "_open_path_with_default_app", lambda _path: None)

    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        wrong_session = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={
                "Authorization": "Bearer secret",
                "x-opensquilla-session-key": "agent:main:webchat:other",
            },
        )
        missing_session = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={"Authorization": "Bearer secret"},
        )

    ArtifactStore(tmp_path).path_for(ref).write_bytes(b"tampered")
    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        mismatch = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={
                "Authorization": "Bearer secret",
                "x-opensquilla-session-key": "agent:main:webchat:ok",
            },
        )

    assert wrong_session.status_code == 404
    assert missing_session.status_code == 404
    assert mismatch.status_code == 409
    assert mismatch.json()["code"] == "INTEGRITY_ERROR"


def test_artifact_native_open_rejects_non_html_artifacts(tmp_path: Path, monkeypatch) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    from openstarry_code.gateway import artifacts as artifact_routes

    ref = _publish(tmp_path, payload=b"plain", name="notes.txt", mime="text/plain")
    opened: list[Path] = []
    monkeypatch.setattr(
        artifact_routes,
        "_open_path_with_default_app",
        lambda path: opened.append(Path(path)) or None,
    )

    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        response = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={
                "Authorization": "Bearer secret",
                "x-opensquilla-session-key": "agent:main:webchat:ok",
            },
        )

    assert response.status_code == 415
    assert response.json()["code"] == "UNSUPPORTED_ARTIFACT_OPEN"
    assert opened == []


def test_artifact_native_open_opener_failure_does_not_leak_path(
    tmp_path: Path,
    monkeypatch,
) -> None:
    pytest.importorskip("starlette.testclient")
    from starlette.testclient import TestClient

    from openstarry_code.gateway import artifacts as artifact_routes

    ref = _publish(tmp_path, payload=b"<html></html>", name="page.html", mime="text/html")
    opened: list[Path] = []

    def fail_open(path: Path) -> str:
        opened.append(Path(path))
        return f"failed to open {path}"

    monkeypatch.setattr(artifact_routes.tempfile, "gettempdir", lambda: str(tmp_path / "tmp"))
    monkeypatch.setattr(artifact_routes, "_open_path_with_default_app", fail_open)

    with TestClient(_app(tmp_path), client=("127.0.0.1", 50000)) as client:
        response = client.post(
            f"/api/v1/artifacts/{ref.id}/open",
            headers={
                "Authorization": "Bearer secret",
                "x-opensquilla-session-key": "agent:main:webchat:ok",
            },
        )

    assert response.status_code == 503
    assert response.json()["code"] == "OPEN_FAILED"
    assert opened
    assert str(opened[0]) not in response.text
